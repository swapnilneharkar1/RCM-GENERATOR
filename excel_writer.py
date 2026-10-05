"""
Writes the assembled RCM rows to a formatted .xlsx workbook, plus a second
sheet for the gap-check QA results.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("process_broad_area", "Process (Broad Area)", 20),
    ("sub_process", "Sub Process (Specific Area)", 24),
    ("stage", "Stage (Stage of Loan Lifecycle)", 20),
    ("risk_addressed", "Risk Addressed", 32),
    ("control_objective", "Control Objective", 28),
    ("control_description", "Control Description", 40),
    ("control_activities", "Control Activities", 32),
    ("mitigating_control", "Mitigating Control", 24),
    ("type_of_control", "Type of Control", 16),
    ("control_classification", "Control Classification", 18),
    ("nature_of_control", "Nature of Control", 14),
    ("control_frequency", "Control Frequency", 14),
    ("control_owner", "Control Owner", 16),
    ("control_reviewer", "Control Reviewer", 16),
    ("information_processing", "Information Processing", 20),
    ("fs_assertion", "Financial Statement Assertion", 24),
    ("key_classification", "Key / Non-Key", 14),
]

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
GAP_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
INFERRED_FILL = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
THIN_BORDER = Border(*(Side(style="thin", color="D9D9D9"),) * 4)
WRAP = Alignment(wrap_text=True, vertical="top")


def write_rcm_workbook(rows: list, gap_findings: list, output_path: str):
    wb = Workbook()

    # --- RCM sheet ---
    ws = wb.active
    ws.title = "RCM"

    for col_idx, (key, label, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"

    for row_idx, row in enumerate(rows, start=2):
        control_desc = str(row.get("control_description", ""))
        is_gap = "control gap identified" in control_desc.lower()
        is_inferred = "[inferred control" in control_desc.lower()

        for col_idx, (key, label, width) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=row.get(key, ""))
            cell.alignment = WRAP
            cell.border = THIN_BORDER
            if is_gap:
                cell.fill = GAP_FILL
            elif is_inferred:
                cell.fill = INFERRED_FILL

    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"

    # --- Legend ---
    legend_row = len(rows) + 3
    ws.cell(row=legend_row, column=1, value="Legend:").font = Font(bold=True)
    ws.cell(row=legend_row + 1, column=1, value="Control Gap Identified").fill = GAP_FILL
    ws.cell(row=legend_row + 2, column=1, value="Inferred Control Based on Documentation").fill = INFERRED_FILL

    # --- Gap-check QA sheet ---
    ws2 = wb.create_sheet("QA - Gap Check")
    headers = ["Checklist Item", "Gap Found", "Detail"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws2.cell(row=1, column=col_idx, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws2.column_dimensions["A"].width = 45
    ws2.column_dimensions["B"].width = 12
    ws2.column_dimensions["C"].width = 80

    for row_idx, item in enumerate(gap_findings, start=2):
        gap_found = item.get("gap_found", False)
        ws2.cell(row=row_idx, column=1, value=item.get("checklist_item", "")).alignment = WRAP
        c2 = ws2.cell(row=row_idx, column=2, value="YES" if gap_found else "No")
        c3 = ws2.cell(row=row_idx, column=3, value=item.get("detail", ""))
        c3.alignment = WRAP
        if gap_found:
            c2.fill = GAP_FILL
            c3.fill = GAP_FILL

    wb.save(output_path)
