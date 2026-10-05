"""
Loads source documents from SOURCE_DOCS_DIR and provides a per-segment
context string.

Supported file types:
- .txt / .md          -> read directly (use this for video TRANSCRIPTS -
                          transcribe audio with Whisper or similar first,
                          then drop the .txt here)
- .pdf                 -> text extracted via pypdf
- .docx                -> text extracted via python-docx
- .pptx                -> slide text + speaker notes
- .xlsx                -> handled SEPARATELY via load_all_rcm_excels() below,
                          not through the generic LOADERS dict. An existing
                          RCM is tabular (1000+ rows) - dumping it as raw
                          text would blow the prompt budget and read as
                          garbled cell-by-cell noise. Instead it's parsed
                          into structured rows and filtered per-segment,
                          same spirit as select_relevant_docs but row-level
                          instead of file-level.
- images (.png/.jpg)   -> NOT read here. If you have walkthrough screenshots,
                          either (a) describe them in a .txt file alongside,
                          or (b) extend `load_image_as_context` below to call
                          a vision-capable model and paste the description in.

Why "select_relevant_docs" instead of dumping everything into every call:
- Keeps each segment's prompt focused (better output quality) and keeps
  token/cost usage down - a 3000-line PPG doesn't need to be re-sent in full
  for the Collections segment if only 2 pages of it are relevant.
- It's a simple keyword filter, not semantic search. For a first prototype
  this is good enough given a handful of documents. If your document set
  grows large, swap this for embeddings-based retrieval (see NOTES.md).
"""

import os
from pathlib import Path

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None


def _read_txt(path: Path) -> str:
    return path.read_text(errors="ignore")


def _read_pdf(path: Path) -> str:
    if PdfReader is None:
        return f"[pypdf not installed - could not read {path.name}]"
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _read_docx(path: Path) -> str:
    if DocxDocument is None:
        return f"[python-docx not installed - could not read {path.name}]"
    doc = DocxDocument(str(path))
    return "\n".join(p.text for p in doc.paragraphs)


def _read_pptx(path: Path) -> str:
    """
    Extracts slide text (titles, bullets, any text boxes) plus speaker
    notes. Does NOT extract or describe embedded images/screenshots pasted
    onto slides - if your PPTX is mostly pasted screenshots with little
    typed text (e.g. sourcing tool screenshots), this will return very
    little useful content. In that case, either type up a short summary of
    each screenshot into a .txt file alongside it, or use a vision model to
    describe each slide image and save that description as text.
    """
    if Presentation is None:
        return f"[python-pptx not installed - could not read {path.name}]"
    prs = Presentation(str(path))
    chunks = []
    for i, slide in enumerate(prs.slides, start=1):
        slide_text = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    slide_text.append(text)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        block = f"[Slide {i}]\n" + "\n".join(slide_text)
        if notes:
            block += f"\n[Speaker notes: {notes}]"
        chunks.append(block)
    return "\n\n".join(chunks)


LOADERS = {
    ".txt": _read_txt,
    ".md": _read_txt,
    ".pdf": _read_pdf,
    ".docx": _read_docx,
    ".pptx": _read_pptx,
}


def load_all_documents(source_dir: str) -> dict:
    """Returns {filename: full_text} for every readable file in source_dir.
    Does NOT include .xlsx files - those go through load_all_rcm_excels()
    instead, since they need row-level (not whole-file) handling."""
    docs = {}
    source_path = Path(source_dir)
    if not source_path.exists():
        return docs
    for file in sorted(source_path.iterdir()):
        if file.suffix.lower() in LOADERS:
            try:
                docs[file.name] = LOADERS[file.suffix.lower()](file)
            except Exception as e:
                docs[file.name] = f"[Error reading {file.name}: {e}]"
    return docs


# ---------- RCM Excel (existing RCM used as PRIMARY generation input) ----------

# Header text (lowercased) -> the field name we store it under. Matched by
# substring, so this tolerates header variations like "Control Description"
# vs "Control  Description\n(...)" across different RCM templates.
FIELD_HEADER_HINTS = [
    ("area", "area"),
    ("business", "business"),
    ("sub process", "subprocess"),
    ("subprocess", "subprocess"),
    ("process", "process"),  # checked after "sub process" so it doesn't win first
    ("control ref", "ctrl_ref"),
    ("stage", "stage"),
    ("products covered", "products"),
    ("risks addressed", "risk"),
    ("risk addressed", "risk"),
    ("control objective", "objective"),
    ("control description", "description"),
    ("control frequency", "frequency"),
    ("control activities", "activities"),
    ("mitigating control", "mitigating"),
    ("type of control", "type"),
    ("control classification", "classification"),
    ("nature of control", "nature"),
    ("criticality", "criticality"),
]


def _map_headers(ws, header_row: int, max_col: int) -> dict:
    col_map = {}
    for c in range(1, max_col + 1):
        h = ws.cell(row=header_row, column=c).value
        if not h:
            continue
        h_lower = str(h).lower()
        for hint, field in FIELD_HEADER_HINTS:
            if hint in h_lower and field not in col_map:
                col_map[field] = c
                break
    return col_map


def load_rcm_excel_rows(path: Path) -> list:
    """
    Parses an existing RCM .xlsx into a list of row-dicts with normalized
    field names (area, process, subprocess, risk, description, etc.).
    Auto-detects the sheet (prefers one named like 'RCM') and the header
    row (scans the first 10 rows for one containing both 'risk' and
    'control' text, since real-world RCM templates often have a title/
    logo above the actual header row).
    """
    if load_workbook is None:
        return []
    wb = load_workbook(str(path), data_only=True)

    ws = None
    for name in wb.sheetnames:
        if "rcm" in name.lower():
            ws = wb[name]
            break
    if ws is None:
        ws = wb[wb.sheetnames[0]]

    header_row = None
    for r in range(1, 11):
        cells = [ws.cell(row=r, column=c).value for c in range(1, min(ws.max_column, 30) + 1)]
        non_empty = [c for c in cells if c]
        row_text = " ".join(str(c) for c in non_empty).lower()
        # A real header row has many distinct populated cells (one per
        # column) AND mentions both risk and control - a title row like
        # "Risk Control Metrix" matches the text check but has only 1-2
        # populated cells, so the count threshold filters it out.
        if len(non_empty) >= 8 and "risk" in row_text and "control" in row_text:
            header_row = r
            break
    if header_row is None:
        header_row = 1

    col_map = _map_headers(ws, header_row, ws.max_column)
    if "risk" not in col_map or "description" not in col_map:
        return []  # doesn't look like an RCM-shaped sheet, skip it

    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        risk_val = ws.cell(row=r, column=col_map["risk"]).value
        desc_val = ws.cell(row=r, column=col_map["description"]).value
        if not risk_val and not desc_val:
            continue
        row = {}
        for field, c in col_map.items():
            v = ws.cell(row=r, column=c).value
            row[field] = str(v).strip() if v is not None else ""
        rows.append(row)
    return rows


def load_all_rcm_excels(source_dir: str) -> list:
    """Scans source_dir for .xlsx files and parses each as RCM reference
    data. Returns a combined list of row-dicts across all such files."""
    all_rows = []
    source_path = Path(source_dir)
    if not source_path.exists():
        return all_rows
    for file in sorted(source_path.iterdir()):
        if file.suffix.lower() == ".xlsx":
            try:
                rows = load_rcm_excel_rows(file)
                all_rows.extend(rows)
            except Exception as e:
                print(f"  [warning] Could not parse {file.name} as an RCM: {e}")
    return all_rows


def filter_relevant_rcm_rows(rcm_rows: list, doc_hints: list) -> list:
    """Returns the matching row-dicts for this segment, UNCAPPED - the
    caller decides how to batch them (see format_rcm_rows_as_text)."""
    if not rcm_rows:
        return []
    relevant = []
    for row in rcm_rows:
        text = " ".join(row.get(k, "") for k in ("area", "process", "subprocess")).lower()
        if any(hint.lower() in text for hint in doc_hints):
            relevant.append(row)
    return relevant


def format_rcm_rows_as_text(rows: list) -> str:
    """Formats a list of matched row-dicts into the compact text block used
    in the prompt. Call this per-batch (see main.py's batching logic) rather
    than on a huge unbatched list, so a single API call's output doesn't
    need to cover more existing rows than it can realistically produce in
    one response."""
    if not rows:
        return ""
    lines = []
    for row in rows:
        lines.append(
            f"- [Ctrl Ref {row.get('ctrl_ref','')}] Area: {row.get('area','')} | "
            f"Process: {row.get('process','')} | Sub-Process: {row.get('subprocess','')}\n"
            f"  Risk: {row.get('risk','')}\n"
            f"  Existing Control Description: {row.get('description','')}\n"
            f"  Existing Control Objective: {row.get('objective','')}\n"
            f"  Type: {row.get('type','')} | Classification: {row.get('classification','')} | "
            f"Nature: {row.get('nature','')} | Frequency: {row.get('frequency','')} | "
            f"Criticality: {row.get('criticality','')}"
        )
    return "\n\n".join(lines)


def select_relevant_rcm_rows(rcm_rows: list, doc_hints: list, max_rows: int = 250) -> str:
    """
    Convenience wrapper kept for callers that just want a single formatted
    block (e.g. quick inspection/testing) - caps at max_rows and truncates.
    main.py's actual generation path uses filter_relevant_rcm_rows +
    format_rcm_rows_as_text directly so it can batch instead of truncating.
    """
    relevant = filter_relevant_rcm_rows(rcm_rows, doc_hints)
    truncated_note = ""
    if len(relevant) > max_rows:
        truncated_note = (f"\n[...{len(relevant) - max_rows} additional existing rows matched "
                           f"this segment but were truncated...]")
        relevant = relevant[:max_rows]
    text = format_rcm_rows_as_text(relevant)
    return text + truncated_note if text else ""


def select_relevant_docs(all_docs: dict, doc_hints: list, max_chars_per_doc: int = 12000) -> str:
    """
    Filters all_docs down to those whose filename matches any doc_hint
    keyword, then concatenates into a single context block. Truncates very
    long docs per-file to keep the prompt manageable (adjust as needed for
    your model's context window).
    """
    relevant = {}
    for fname, text in all_docs.items():
        fname_lower = fname.lower()
        if any(hint.lower() in fname_lower for hint in doc_hints):
            relevant[fname] = text

    # Fallback: if nothing matched by filename, include everything
    # (better to over-include than silently produce an empty-context call).
    if not relevant:
        relevant = all_docs

    blocks = []
    for fname, text in relevant.items():
        truncated = text[:max_chars_per_doc]
        note = "" if len(text) <= max_chars_per_doc else "\n[...truncated...]"
        blocks.append(f"--- DOCUMENT: {fname} ---\n{truncated}{note}\n")

    return "\n".join(blocks) if blocks else "[No source documents available for this segment - rely on the INFERENCE RULE and general NBFC-MFI / RBI knowledge, clearly marking all such rows as inferred.]"
